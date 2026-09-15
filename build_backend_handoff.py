"""Package truthful missing values for backend import and a no-blank review view."""
import hashlib
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from prepare_agent_data import META, missing, sheet

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'output-handoff'
NUMERIC = {'rent_monthly_sgd', 'bedrooms', 'bathrooms', 'area_sqm',
           'mrt_distance_m', 'mrt_minutes_reported', 'tenant_count_reported'}
DATES = {'listed_date', 'collected_at'}
REQUIRED = {'listing_id', 'title', 'listing_url', 'rent_monthly_sgd', 'district_code', 'address'}


def main():
    source = ROOT / 'output-agent/agent_listings.json'
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    rows = json.loads(source.read_text(encoding='utf-8'))
    fields = list(rows[0])
    assert len(rows) == 600 and len(fields) == 33
    assert len({r['listing_id'] for r in rows}) == len(rows)
    issues = []
    for r in rows:
        assert set(r) == set(fields)
        for f in REQUIRED:
            if missing(r[f]):
                issues.append({'listing_id': r['listing_id'], 'field': f, 'issue': 'missing_required_value'})
        for f in NUMERIC:
            assert r[f] is None or (isinstance(r[f], (int, float)) and not isinstance(r[f], bool)), (r['listing_id'], f)
    assert not issues, issues
    metadata, missing_by_listing = [], []
    for f in fields:
        null_count = sum(r[f] is None for r in rows)
        unknown_count = sum(r[f] == 'unknown' for r in rows)
        metadata.append(dict(field=f, meaning_zh=META[f]['meaning_zh'],
                             value_type='number' if f in NUMERIC else 'date_string' if f in DATES else 'string_array' if f == 'review_flags' else 'string',
                             required_for_ingestion=f in REQUIRED,
                             missing_representation='null' if f in NUMERIC | DATES else 'empty_array_means_no_flags' if f == 'review_flags' else 'unknown',
                             null_count=null_count, unknown_count=unknown_count,
                             known_count=len(rows)-null_count-unknown_count))
    for r in rows:
        missing_by_listing.append(dict(listing_id=r['listing_id'],
                                       missing_fields=[f for f in fields if missing(r[f])],
                                       reason='missing_or_unconfirmed_in_prepared_data'))
    review = [{f: 'unknown' if v is None or v == '' else v for f, v in r.items()} for r in rows]
    null_total = sum(m['null_count'] for m in metadata)
    summary = dict(rows=len(rows), fields=len(fields), missing_required_values=0,
                   numeric_or_date_null_cells=null_total,
                   categorical_unknown_cells=sum(m['unknown_count'] for m in metadata),
                   review_blanks_replaced_with_unknown=null_total,
                   factual_values_imputed=0, records_deleted=0, source_sha256=source_hash,
                   policy='Typed import keeps null; review display uses unknown. Neither creates new factual information.')
    OUT.mkdir(exist_ok=True)
    for name, value in [('listings_for_backend.json', rows), ('field_contract.json', metadata),
                        ('missing_fields_by_listing.json', missing_by_listing), ('handoff_report.json', summary)]:
        (OUT / name).write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')
    wb = Workbook()
    wb.remove(wb.active)
    sheet(wb, 'Review_No_Blanks', review)
    sheet(wb, 'Field_Contract', metadata)
    sheet(wb, 'Missing_By_Listing', missing_by_listing)
    sheet(wb, 'Read_Me', [
        {'item': 'purpose', 'instruction': 'Human review only. Numeric unknown text must not be imported into numeric database columns.'},
        {'item': 'database_import', 'instruction': 'Import listings_for_backend.json. Preserve JSON null as database NULL; preserve categorical unknown.'},
        {'item': 'meaning', 'instruction': 'unknown and null mean missing or unconfirmed, never zero, no, or unrestricted.'},
        {'item': 'hard_constraints', 'instruction': 'Missing evidence -> needs_confirmation. Never silently pass a hard requirement.'},
        {'item': 'soft_preferences', 'instruction': 'Report evidence coverage separately; do not treat unknown distance as zero or as evidence of poor location.'},
        {'item': 'scope', 'instruction': 'Ingestion checks do not establish live availability or suitability. NUS commute and tenant-count scope remain unverified.'},
    ])
    path = OUT / 'nus_rental_600_review_no_blanks.xlsx'
    wb.save(path)
    (OUT / 'README_BC_ZH.md').write_text(f'''# B/C 数据交付说明

当前保留 {len(rows)} 条、{len(fields)} 字段，没有删除记录或虚构补全事实。

## 文件用途

- `listings_for_backend.json`：B 导入数据库、C/后端读取的数值类型保持正确的数据。可选数值仍有 null，分类未知值为 unknown。
- `nus_rental_600_review_no_blanks.xlsx`：供小组人工查看；原来的 {null_total} 个数值/日期空白显示为 unknown。不能直接按数值表导入数据库。
- `field_contract.json`：每个字段的类型、必填性、缺失规则和中文含义。
- `missing_fields_by_listing.json`：按 listing_id 记录每条缺失/未确认字段，供接口生成待确认提示。
- `handoff_report.json`：处理数量与源文件校验值。

## 实际处理及合理性

1. 核验编号唯一、核心必填字段完整、数值列类型正确。当前必填字段缺失为 0，600 条通过入库基本检查。
2. 将 Excel 查看表中的空白统一显示 unknown，让组员看得出“缺失”，而不是误以为漏写。只是显示变化，没有新增事实。
3. 数据库交付版保持 JSON null，不在数值字段混入 unknown，也不填 0、-1、平均数或其它房源的距离。
4. 分类字段保持 unknown，明确的 yes/no 值不改动。Studio 独立卧室数 0 等真实值仍为 0。
5. 数据源原值、地区比例、价格、房源编号均保持。源文件未覆盖。没有重新爬取或声称补查到了缺失的信息。

## B 的入库规则

必填字段：listing_id、title、listing_url、rent_monthly_sgd、district_code、address。
listing_id 按文本存储，listing_id 做唯一键，listing_url 做唯一性校验；可选数值允许 NULL。
只在入库校验阶段拒绝/隔离必填字段不合格的记录，不因可选字段缺失删除房源。
不要对可选数值设置默认 0。数值排序时将缺失项另列或放到末尾，不能把它们当作最低价、最近或最小面积。
以上仅为入库数据契约，不代表已完成数据库部署；按组内选定数据库实现。

## C / 后端 / Agent 的读取规则

- JSON null、字段缺失、分类 unknown，都表示当前无法确认；数值 0 表示真实数值。
- 硬条件：已知符合 -> pass；已知不符合 -> fail；缺失/冲突 -> needs_confirmation。严格推荐与待确认候选分开展示。
- 软偏好：不要伪造值参与比较。匹配度和信息覆盖率分别呈现；未注明不等于差，也不能因为只剩少数已知字段就表现为高置信度。
- cooking_policy 缺失不能判断能做饭。tenant_count_reported 不是已入住人数，人数标签的适用范围尚未确认。
- MRT 分钟数不是到 NUS 的通勤时间；运营状态 unknown 不是已核实运营。
- 给用户展示时把 null 显示为“未注明”，LLM 不得猜测补全。

## 仍存在的限制

没有空白的查看表不等于信息齐全。若希望获得真实缺失值，需要进一步获取可核验的房源详情或路线数据。
当前交付适合数据库/接口开发与评审，不保证所有用户的硬条件都能被现有证据验证。
''', encoding='utf-8')
    # Meaningful release checks: every original value preserved in import; every
    # review cell equals its original value except the explicit display substitution.
    imported = json.loads((OUT / 'listings_for_backend.json').read_text(encoding='utf-8'))
    assert imported == rows
    check = load_workbook(path, read_only=True)
    ws = check['Review_No_Blanks']
    assert ws.max_row == 601 and ws.max_column == 33
    for cells, original in zip(ws.iter_rows(min_row=2), rows):
        for f, cell in zip(fields, cells):
            expected = original[f]
            if expected is None or expected == '':
                expected = 'unknown'
            elif isinstance(expected, (list, dict)):
                expected = json.dumps(expected, ensure_ascii=False)
            assert cell.value is not None and cell.value == expected, (original['listing_id'], f)
    check.close()
    assert hashlib.sha256(source.read_bytes()).hexdigest() == source_hash
    print(json.dumps(summary, indent=2))


if __name__ == '__main__':
    main()
