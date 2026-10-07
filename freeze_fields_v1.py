"""Freeze a proposed V1 schema using only the existing 33 fields."""
import hashlib
import json
from pathlib import Path
from openpyxl import Workbook, load_workbook
from prepare_agent_data import META, missing, sheet

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'output-final-v1'
AUX = {
    'district': ('District_Lookup', '分区名称按 district_code 查询即可；避免每条重复存储。展示时仍可使用。'),
    'availability_status': ('Listing_Metadata', '缺少可入住的完整日期，不能可靠筛选入住时间；保留快照状态供回查。'),
    'review_flags': ('Listing_Metadata', '属于质量控制而非租房偏好；后端必须关联读取，不能因移出主表忽略警告。'),
}
NUMBERS = {'rent_monthly_sgd', 'bedrooms', 'bathrooms', 'area_sqm', 'mrt_distance_m', 'mrt_minutes_reported', 'tenant_count_reported'}
NOTES = {
    'tenant_count_reported': '仅作广告人数线索；范围未确认，不能直接用于硬性入住人数判断，也不代表已有租客。',
    'tenant_gender_label': '仅为平台标签；不推断现住室友性别，不将标签自动升级为排他限制。',
    'gender_restriction': '明确限制才能用于兼容性判断；仅对用户主动提供的相关需求进行判断。',
    'gender_preference': '广告偏好，不是排他限制；不得当作硬性拒绝规则。',
    'area_scope': '必须与 area_sqm 一起使用，避免比较房间面积与整套或土地面积。',
    'mrt_operation_status': 'under_construction 是明确原文标记；unknown 不表示已运营。',
    'mrt_minutes_reported': '广告到车站的时间，不是到 NUS 的通勤时间；交通方式尚未核实。',
    'rental_scope': '未知不能自动当作整租；只依据明确文字或房间标签。',
    'bedrooms': '不是室友数或床位数；Studio 的 0 指零个独立卧室；房间广告的数量范围需要核实。',
    'bathrooms': '不是独享卫生间数量；独享性需参考 private_bathroom。',
    'rent_price_type': '起价/指导价不能保证最终月租；asking 仅为广告报价。',
    'listed_date': '广告发布日期，不代表仍然有效或当前可租。',
    'collected_at': '数据采集时间，不是房源更新时间或最后核实可租时间。',
}


def main():
    src = ROOT / 'output-agent/agent_listings.json'
    before = hashlib.sha256(src.read_bytes()).hexdigest()
    original = json.loads(src.read_text(encoding='utf-8'))
    fields = [f for f in original[0] if f not in AUX]
    assert len(fields) == 30 and len(original) == 600
    records = [{f: r[f] for f in fields} for r in original]
    metadata = [{f: r[f] for f in ('listing_id', 'availability_status', 'review_flags')} for r in original]
    districts = {}
    for r in original:
        code, name = r['district_code'], r['district']
        assert code not in districts or districts[code] == name
        districts[code] = name
    lookup = [{'district_code': k, 'district': v} for k, v in sorted(districts.items())]
    decisions = []
    for f in original[0]:
        decisions.append(dict(field=f, meaning_zh=META[f]['meaning_zh'],
                              decision='move_to_auxiliary' if f in AUX else 'keep_in_main',
                              destination=AUX[f][0] if f in AUX else 'Listings',
                              valid_count=sum(not missing(r[f]) and r[f] != [] for r in original),
                              decision_role=META[f]['decision_role'],
                              value_type='number_or_null' if f in NUMBERS else 'array_of_strings' if f == 'review_flags' else 'ISO_date_string' if f in ('listed_date', 'collected_at') else 'string',
                              reason_zh=AUX[f][1] if f in AUX else NOTES.get(f, '用于房源筛选、比较、展示或溯源；缺失率高不等于字段无用，未知不得猜测补全。')))
    OUT.mkdir(exist_ok=True)
    payload = dict(schema_version='proposed_v1', listings=records, listing_metadata=metadata, district_lookup=lookup)
    (OUT / 'rental_dataset_v1.json').write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
    (OUT / 'field_decisions_v1.json').write_text(json.dumps(decisions, ensure_ascii=False, indent=2), encoding='utf-8')
    wb = Workbook()
    wb.remove(wb.active)
    # Explicitly a human-readable view; import typed JSON, not numeric "unknown" text.
    review = [{f: 'unknown' if v is None else v for f, v in r.items()} for r in records]
    sheet(wb, 'Listings_Review', review)
    sheet(wb, 'Final_30_Fields', [r for r in decisions if r['decision'] == 'keep_in_main'])
    sheet(wb, 'All_33_Decisions', decisions)
    sheet(wb, 'Listing_Metadata', metadata)
    sheet(wb, 'District_Lookup', lookup)
    sheet(wb, 'Read_Me', [
        dict(item='status', value='Proposed V1 field freeze for group approval; not a claim that every field is complete.'),
        dict(item='database', value='Import rental_dataset_v1.json. Join listing_metadata by listing_id and district_lookup by district_code.'),
        dict(item='display', value='Listings_Review replaces null with unknown for humans only. JSON retains numeric null.'),
        dict(item='quality', value='Backend must read review_flags from listing_metadata before presenting confirmed recommendations.'),
        dict(item='limits', value='Unknown hard requirements need confirmation. No verified NUS commute, coordinates, full costs or live availability in this schema.'),
    ])
    target = OUT / 'NUS_Rental_V1_Fields_and_Data.xlsx'
    wb.save(target)
    lines = ['# 建议提交小组的 V1 字段方案', '',
             '已有 33 个字段中，30 个进入房源主表，3 个移入辅助表，没有物理删除原始信息。保留全部 600 条。',
             '这是一份建议冻结并提交组员确认的 V1 方案，不代表已经获得小组批准，也不代表已覆盖完整 Product Guide 的全部数据需求。', '',
             '## 给 B 和 C', '',
             '- B：导入 rental_dataset_v1.json，其中 listings 为 30 字段主表；listing_metadata 按 listing_id 关联；district_lookup 按 district_code 关联。',
             '- C：读取房源必须同时读取关联 review_flags；有冲突或相关信息缺失时，返回需要确认，而不是默认为满足。',
             '- Excel 供人工查看，将 null 显示为 unknown；数据库导入使用 JSON，数字缺失仍为 null。',
             '- 字段是否必填、是否硬条件是不同概念；项目默认硬条件为预算、出租范围、卧室需求，但需求值未知时不能通过。',
             '- 物业类型、房型、家具、面积、区域与车站距离通常是软偏好；用户要求“必须”时可变为硬条件。',
             '- cooking_policy、private_bathroom、air_conditioning、landlord_lives_in、pets_policy 等虽覆盖少，但对学生选房有用，保留，不因缺失率高删除。',
             '- 人数标签和性别标签只作为带有不确定性的证据；不要把人数当作已入住人数或自动认定为房间最大容量。',
             '- listed_date / collected_at 用于数据溯源，不保证当前仍可租。', '',
             '## 不进入主表的 3 个字段', '']
    lines += [f'- `{k}` → {v[0]}：{v[1]}' for k, v in AUX.items()]
    lines += ['', '## 主表全部 30 字段', '', '| Field | 中文含义 | 有效数量 | 使用说明 |', '|---|---|---:|---|']
    lines += [f"| {d['field']} | {d['meaning_zh']} | {d['valid_count']} | {d['reason_zh']} |" for d in decisions if d['decision'] == 'keep_in_main']
    lines += ['', '## 边界', '',
              '本轮仅考虑已有字段，不重新加入之前全空的可入住日期、租期、水电等字段，也不新增路线或地理字段。',
              '当前没有可靠坐标、NUS 通勤时间、完整总费用和实时可租状态。不能据此宣称已满足地图、通勤及费用核实的所有需求。',
              '没有删除低覆盖率房源，没有补造缺失事实，没有改变区域配额。所有原值可从主表和辅助表无损还原。']
    (OUT / 'SUBMISSION_GUIDE_ZH.md').write_text('\n'.join(lines), encoding='utf-8')
    # Verify lossless separation, joins, row counts, and every worksheet value.
    saved = json.loads((OUT / 'rental_dataset_v1.json').read_text(encoding='utf-8'))
    meta_by_id = {r['listing_id']: r for r in saved['listing_metadata']}
    for old, r in zip(original, saved['listings']):
        restored = dict(r, **{k: v for k, v in meta_by_id[r['listing_id']].items() if k != 'listing_id'}, district=districts[r['district_code']])
        assert restored == old
    check = load_workbook(target, read_only=True)
    assert check['Listings_Review'].max_row == 601 and check['Listings_Review'].max_column == 30
    assert check['Final_30_Fields'].max_row == 31
    for actual, expected in zip(check['Listings_Review'].iter_rows(min_row=2), review):
        assert [c.value for c in actual] == list(expected.values())
    check.close()
    assert hashlib.sha256(src.read_bytes()).hexdigest() == before
    print(json.dumps(dict(main_fields=len(fields), auxiliary_fields=list(AUX), rows=len(records), districts=len(lookup), lossless_roundtrip=True)))


if __name__ == '__main__':
    main()
