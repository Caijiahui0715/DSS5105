"""Conservative, reproducible preparation of saved rental snapshots for an agent."""
import csv
import hashlib
import json
import math
import re
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from urllib.parse import urlsplit, urlunsplit

from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'output-agent'
UNKNOWN = 'unknown'

# The data fields are English. Chinese is confined to the explanatory dictionary.
DEFINITIONS = '''listing_id|房源编号，文本格式|identity
title|房源名称，保留原文|display
listing_url|具体房源详情链接，不是搜索页|source
rent_monthly_sgd|广告月租，新加坡元；不是每人分摊价或含杂费总价|hard_default
rent_price_type|报价性质；asking/negotiable/starting_from/guide_price/unknown|hard_support
rental_scope|出租范围；entire_unit/room/unknown，必须有明确依据|hard_default
room_type|common_room/master_room/shared_room/studio/room_unspecified/unknown|hard_or_soft
property_type|标准化物业类型，如 condominium、hdb_flat|soft_default
bedrooms|广告卧室数；明确 Studio 记为 0 个独立卧室；不是合租室友数|hard_default
bathrooms|广告卫生间数；未承诺属于出租房间独享|hard_or_soft
area_sqm|展示面积，统一平方米；必须结合 area_scope 使用|soft_default
area_scope|floor_reported/land_reported/unknown；不保证是出租房间面积|quality
district_code|PPG 原始分区 D01–D28；不按地址自行重新分区|hard_or_soft
district|PPG 分区英文名称|display
address|原始地址，整理空白字符，保留门牌|hard_or_soft
mrt_station|广告所列 MRT/LRT 站；不代表已核实运营或最近车站|soft_default
mrt_distance_m|广告到所列车站距离，米；不是到 NUS 距离|soft_default
mrt_minutes_reported|广告到所列车站分钟数；不是 NUS 通勤时间|soft_default
mrt_operation_status|只有原文 U/C 时记 under_construction，否则 unknown|quality
available_from|明确写出四位年份时的可入住日期 YYYY-MM-DD，否则留空|hard_or_soft
availability_status|date_reported/ready_at_snapshot/date_missing_year/unknown|hard_support
listed_date|广告刊登日期 YYYY-MM-DD；不是现在仍可租的证明|source
collected_at|原数据采集时间，统一 ISO 格式；不补猜缺失时区|source
tenant_count_reported|广告人数标签，不能理解为已入住人数或保证可入住人数|hard_support
tenant_count_scope|unknown；现有标签不足以可靠区分房间或整套容量|quality
tenant_gender_label|平台显示的 male/female/unknown 标签；不代表当前室友性别|hard_support
gender_restriction|只有明确 only 等排他措辞才提取 male_only/female_only|hard_or_soft
gender_preference|只有明确 preferred 措辞才提取 male/female；不是硬性拒绝条件|soft_default
cooking_policy|allowed/light_only/not_allowed/unknown；不写做饭就是 unknown|hard_or_soft
furnishing|fully_furnished/partially_furnished/furnished_unspecified/unfurnished/unknown|soft_default
landlord_lives_in|yes/no/unknown；no owner 指广告声称房东不住，不代表没有业主|hard_or_soft
private_bathroom|yes/no/unknown；主卧不自动意味着独立卫浴|hard_or_soft
air_conditioning|yes/no/unknown；仅依据明确广告文字|hard_or_soft
utilities_included|yes/no/unknown；仅提到 utilities 不等于费用包含|hard_or_soft
wifi_included|yes/no/unknown；仅提到 wifi 不等于费用包含|hard_or_soft
visitors_policy|allowed/not_allowed/unknown|hard_or_soft
pets_policy|allowed/not_allowed/unknown|hard_or_soft
smoking_policy|allowed/not_allowed/unknown|hard_or_soft
minimum_lease_months|明确最低租期，月；preferred lease 不作为最低租期|hard_or_soft
review_flags|待核对标记列表，不自动删房源|quality'''
META = {a: {'field': a, 'meaning_zh': b, 'decision_role': c}
        for a, b, c in (line.split('|') for line in DEFINITIONS.splitlines())}
RESTRICTIONS = ['gender_restriction', 'gender_preference', 'cooking_policy', 'furnishing',
                'landlord_lives_in', 'private_bathroom', 'air_conditioning', 'utilities_included',
                'wifi_included', 'visitors_policy', 'pets_policy', 'smoking_policy', 'minimum_lease_months']


def missing(v):
    return v is None or (isinstance(v, str) and v.strip().lower() in ('', 'unknown', 'null', 'none', 'nan'))


def text(v):
    return UNKNOWN if missing(v) else re.sub(r'\s+', ' ', str(v)).strip()


def number(v):
    if missing(v) or isinstance(v, bool):
        return None
    try:
        n = float(str(v).replace(',', '').replace('S$', '').strip())
        return n if math.isfinite(n) else None
    except ValueError:
        return None


def snake(v):
    return re.sub(r'[^a-z0-9]+', '_', text(v).lower()).strip('_')


def date_iso(v, full=False):
    if missing(v):
        return None
    try:
        d = datetime.fromisoformat(str(v).replace('Z', '+00:00'))
        return d.isoformat(timespec='seconds') if full else d.date().isoformat()
    except ValueError:
        for fmt in ('%b %d, %Y', '%d %b %Y', '%d %B %Y'):
            try:
                return datetime.strptime(str(v), fmt).date().isoformat()
            except ValueError:
                pass
    return None


def canonical_url(v):
    p = urlsplit(str(v or '').strip())
    return urlunsplit((p.scheme.lower(), p.netloc.lower(), p.path.rstrip('/'), '', ''))


def rendered_cards(ids):
    from page_dom import page_info
    cards = defaultdict(list)
    for path in sorted((ROOT / 'saved-verified').glob('*.html')):
        soup = BeautifulSoup(path.read_text(encoding='utf-8'), 'html.parser')
        source = page_info(soup)['source_url']
        for card in soup.select('.listing-card-v2[da-listing-id]'):
            key = card['da-listing-id']
            if key not in ids:
                continue
            values = {}
            for suffix in ('headline', 'bedrooms', 'bathrooms', 'room-type', 'area', 'availability'):
                node = card.find(attrs={'da-id': 'listing-card-v2-' + suffix})
                values[suffix] = node.get_text(' ', strip=True) if node else ''
            values.update(file=path.name, source_url=source)
            cards[key].append(values)
    return cards


def extract_restrictions(description):
    """Rules operate on the saved card headline, not agents' names or page-wide text."""
    result = {f: UNKNOWN for f in RESTRICTIONS}
    result['minimum_lease_months'] = None
    evidence, flags = {}, []
    rules = {
        'cooking_policy': {
            'not_allowed': r'\b(?:no cooking|cooking (?:is )?not allowed)\b',
            'light_only': r'\b(?:light|simple) cooking(?: only| allowed)?\b',
            'allowed': r'(?<!light )(?<!simple )\bcooking (?:is )?(?:allowed|permitted)\b'},
        'furnishing': {
            'fully_furnished': r'\bfully furnish(?:ed)?\b|\bfully furn\b',
            'partially_furnished': r'\bpartial(?:ly)? furnish(?:ed)?\b',
            'unfurnished': r'\bunfurnished\b'},
        'landlord_lives_in': {
            'no': r'\bno (?:live[- ]in )?(?:owner|landlord)\b',
            'yes': r'(?<!no )\b(?:owner|landlord) (?:lives|stays|staying) (?:in|with)\b'},
        'private_bathroom': {
            'yes': r'\b(?:private|own|attached|en[- ]?suite) (?:bathroom|bath)\b',
            'no': r'\bshared bathroom\b'},
        'air_conditioning': {'yes': r'(?<!no )\baircon\b|\bair[- ]conditioned\b', 'no': r'\bno aircon\b'},
        'utilities_included': {'yes': r'\b(?:utilities|pub) (?:are )?included\b', 'no': r'\b(?:utilities|pub) (?:are )?(?:excluded|not included)\b'},
        'wifi_included': {'yes': r'\bwi[- ]?fi (?:is )?included\b', 'no': r'\bwi[- ]?fi (?:is )?(?:excluded|not included)\b'},
        'visitors_policy': {'allowed': r'\bvisitors (?:are )?allowed\b', 'not_allowed': r'\bno visitors\b|\bvisitors (?:are )?not allowed\b'},
        'pets_policy': {'allowed': r'\bpet[- ]friendly\b|\bpets (?:are )?allowed\b', 'not_allowed': r'\bno pets\b|\bpets (?:are )?not allowed\b'},
        'smoking_policy': {'allowed': r'\bsmoking (?:is )?allowed\b', 'not_allowed': r'\bno smoking\b|\bsmoking (?:is )?not allowed\b'},
        'gender_restriction': {'female_only': r'\bfemale only\b|\bonly (?:single )?(?:lad(?:y|ies)|female)\b|\ball ladies\b',
                               'male_only': r'\bmale only\b|\bonly (?:single )?male\b'},
        'gender_preference': {'male': r'\bmale (?:tenant )?preferred\b', 'female': r'\bfemale (?:tenant )?preferred\b'},
    }
    for field, patterns in rules.items():
        found = []
        for value, pat in patterns.items():
            for m in re.finditer(pat, description, re.I):
                # Unhandled negation/conditional wording is not affirmative evidence.
                prefix = description[max(0, m.start()-20):m.start()]
                if re.search(r'\b(?:not|no|if|possibly|maybe)\s*$', prefix, re.I):
                    flags.append('ambiguous_' + field)
                    continue
                found.append((value, m.group()))
        values = {v for v, _ in found}
        if len(values) == 1:
            result[field] = next(iter(values))
        elif len(values) > 1:
            flags.append('conflicting_' + field)
        if found:
            evidence[field] = [{'value': v, 'quote': q} for v, q in found]
    if result['furnishing'] == UNKNOWN and not evidence.get('furnishing') and 'ambiguous_furnishing' not in flags:
        m = re.search(r'\bfurnished\b', description, re.I)
        if m and not re.search(r'\b(?:not|no)\s*$', description[:m.start()], re.I):
            result['furnishing'] = 'furnished_unspecified'
            evidence['furnishing'] = [{'value': result['furnishing'], 'quote': m.group()}]
    m = re.search(r'\b(?:minimum|min\.?|at least)\s+(\d+)\s*[- ]?\s*(months?|years?)\b', description, re.I)
    if m:
        result['minimum_lease_months'] = int(m[1]) * (12 if m[2].lower().startswith('year') else 1)
        evidence['minimum_lease_months'] = [{'value': result['minimum_lease_months'], 'quote': m[0]}]
    return result, evidence, flags


def clean_record(raw, observations):
    flags, transformations = [], []
    matches = [o for o in observations if o['source_url'] == raw['source_url']]
    selected = matches or observations
    # Only use a single consistent saved feature value; never choose silently between conflicts.
    dom = {}
    for k in ('headline', 'bedrooms', 'bathrooms', 'room-type', 'area', 'availability'):
        values = {o[k] for o in selected}
        dom[k] = next(iter(values)) if len(values) == 1 else ''
        if len(values) > 1:
            flags.append('conflicting_saved_' + k.replace('-', '_'))
    if not selected:
        flags.append('saved_card_not_found')
    headline = dom['headline']
    rules, evidence, rule_flags = extract_restrictions(headline)
    flags += rule_flags
    scope = 'room' if snake(raw.get('rental_scope')) == 'room' else UNKNOWN
    entire = re.search(r'\b(?:whole|entire) (?:unit|apartment|house)\b', headline, re.I)
    room = re.search(r'\b(?:common|master|shared) (?:room|bedroom)\b', headline, re.I)
    if entire:
        if scope == 'room' or room:
            scope = UNKNOWN
            flags.append('conflicting_rental_scope')
        else:
            scope = 'entire_unit'
        evidence['rental_scope'] = entire[0]
    elif room:
        scope = 'room'
        evidence['rental_scope'] = room[0]
    room_type = snake(raw.get('room_type'))
    if room_type == UNKNOWN and room:
        room_type = room[0].lower().split()[0] + '_room'
    beds = number(raw.get('bedrooms'))
    if dom['bedrooms'].lower() == 'studio':
        if room_type not in (UNKNOWN, 'studio'):
            flags.append('studio_room_type_conflict')
        else:
            room_type = 'studio'
        beds = 0
        transformations.append('explicit_studio_to_zero_independent_bedrooms')
    if beds is not None and beds < 0:
        beds = None
        transformations.append('negative_bedroom_sentinel_to_null')
    if room_type == UNKNOWN and scope == 'room':
        room_type = 'room_unspecified'
    baths = number(raw.get('bathrooms'))
    if baths is not None and baths <= 0:
        baths = None
        transformations.append('nonpositive_bathroom_to_null')
    floor, land, legacy = [number(raw.get(k)) for k in ('floor_area_sqft', 'land_area_sqft', 'area_sqft')]
    sqft = floor if floor is not None else land if land is not None else legacy
    area_scope = 'floor_reported' if floor is not None else 'land_reported' if land is not None else UNKNOWN
    if scope == 'room':
        area_scope = UNKNOWN
    if sqft is not None and sqft <= 0:
        sqft = None
        flags.append('invalid_area_to_null')
    price = number(raw.get('price_sgd'))
    if text(raw.get('currency')) != 'SGD' or text(raw.get('rent_period')) != 'month':
        flags.append('unconfirmed_monthly_sgd_basis')
        price = None
    if price is not None and price <= 0:
        flags.append('invalid_rent_to_null')
        price = None
    if price is not None and not 300 <= price <= 30000:
        flags.append('rent_outside_review_range')
    if sqft is not None and not 80 <= sqft <= 10000:
        flags.append('area_outside_review_range')
    if scope == 'room' and beds is not None and beds > 1:
        flags.append('room_bedroom_scope_review')
    psf = number(raw.get('psf_sgd'))
    comparable_areas = [a for a in (floor, land) if a and a > 0]
    if scope != 'room' and psf and price and comparable_areas:
        if all(abs(price / a - psf) / psf > 0.15 for a in comparable_areas):
            flags.append('price_area_psf_mismatch')
    original_sqm = number(raw.get('floor_area_sqm'))
    if floor and original_sqm and abs(floor * 0.09290304 - original_sqm) > max(1, original_sqm * 0.02):
        flags.append('original_area_unit_mismatch')
    availability = text(raw.get('availability_text'))
    available_from, availability_status = None, UNKNOWN
    if re.search(r'\bready to move\b', availability, re.I):
        availability_status = 'ready_at_snapshot'
    elif availability != UNKNOWN:
        candidate = re.sub(r'^Available from\s*', '', availability, flags=re.I)
        available_from = date_iso(candidate)
        availability_status = 'date_reported' if available_from else 'date_missing_year' if not re.search(r'\b\d{4}\b', candidate) else UNKNOWN
    label = text(raw.get('tenant_gender_label')).lower()
    count = number(raw.get('tenant_count_reported'))
    if count is not None:
        flags.append('tenant_count_scope_unconfirmed')
        m = re.search(r'\b(\d+)\s*[-–]\s*(\d+)\s*pax\b', headline, re.I)
        if m and not int(m[1]) <= count <= int(m[2]):
            flags.append('headline_tenant_count_conflict')
    if rules['gender_restriction'] != UNKNOWN and label != UNKNOWN and not rules['gender_restriction'].startswith(label):
        flags.append('gender_label_text_conflict')
        rules['gender_restriction'] = UNKNOWN
    url = canonical_url(raw.get('url'))
    p = urlsplit(url)
    if p.scheme != 'https' or p.hostname != 'www.propertyguru.com.sg' or not p.path.startswith('/listing/') or not p.path.endswith('-' + str(raw['listing_id'])):
        flags.append('invalid_listing_url_or_id')
    # Previous checks remain traceable, including conflicting PSF and area units.
    old_flags = raw.get('review_flags')
    if isinstance(old_flags, str):
        try:
            old_flags = json.loads(old_flags)
        except ValueError:
            old_flags = [x.strip() for x in old_flags.split(';') if x.strip()]
    flags += ['previous_' + f for f in (old_flags or []) if isinstance(f, str)
              and f not in ('negative_numeric', 'area_review', 'price_review', 'same_address_different_ids')]
    price_type = snake(raw.get('price_terms'))
    if price_type == UNKNOWN and price is not None:
        price_type = 'asking'
        transformations.append('unqualified_advertised_monthly_price_to_asking')
    row = dict(listing_id=str(raw['listing_id']), title=text(raw.get('title')), listing_url=url,
               rent_monthly_sgd=price, rent_price_type=price_type,
               rental_scope=scope, room_type=room_type, property_type=snake(raw.get('property_type')),
               bedrooms=beds, bathrooms=baths, area_sqm=round(sqft * 0.09290304, 2) if sqft else None,
               area_scope=area_scope, district_code=text(raw.get('district_code')), district=text(raw.get('district')),
               address=text(raw.get('address')), mrt_station=text(raw.get('mrt_station')),
               mrt_distance_m=number(raw.get('mrt_distance_m')), mrt_minutes_reported=number(raw.get('mrt_minutes_reported')),
               mrt_operation_status='under_construction' if '(U/C)' in str(raw.get('mrt_text')) else UNKNOWN,
               available_from=available_from, availability_status=availability_status,
               listed_date=date_iso(raw.get('posted_at_text')), collected_at=date_iso(raw.get('collected_at'), True),
               tenant_count_reported=count, tenant_count_scope=UNKNOWN, tenant_gender_label=label,
               **rules, review_flags=sorted(set(flags)))
    provenance = dict(listing_id=row['listing_id'], source_url=raw['source_url'], sample_cohort=raw.get('sample_cohort'),
                      source_files=sorted({o['file'] for o in selected}), headline=headline,
                      availability_text=raw.get('availability_text'), tenant_label_text=raw.get('tenant_label_text'),
                      card_text=raw.get('card_text'), thumbnail_url=raw.get('thumbnail_url'), image_urls=raw.get('image_urls'),
                      original_field_evidence=raw.get('field_evidence'), original_review_flags=old_flags,
                      rule_evidence=evidence, transformations=transformations)
    return row, provenance


def assess_constraints(row, requirements):
    """Tri-state checks. Unknown hard requirements never pass silently.

    Supported requirement keys are the same as field names except max_rent_monthly_sgd
    and min_bedrooms. Unsupported keys require confirmation (e.g. NUS commute).
    """
    checks = []
    flags = row.get('review_flags', [])
    for key, wanted in requirements.items():
        field = {'max_rent_monthly_sgd': 'rent_monthly_sgd', 'min_bedrooms': 'bedrooms'}.get(key, key)
        value = row.get(field)
        state, reason = 'needs_confirmation', 'missing_or_unverified'
        if not missing(value):
            if key == 'max_rent_monthly_sgd':
                if row.get('rent_price_type') in ('starting_from', 'guide_price', UNKNOWN, None):
                    reason = 'price_basis_requires_confirmation'
                else:
                    state = 'pass' if value <= wanted else 'fail'
                    reason = 'advertised_rent_comparison_excludes_extra_costs'
            elif key == 'min_bedrooms':
                if row.get('rental_scope') == 'room':
                    reason = 'room_and_dwelling_bedroom_scope_differ'
                else:
                    state = 'pass' if value >= wanted else 'fail'
                    reason = 'bedroom_comparison'
            elif field in META and META[field]['decision_role'] in ('hard_default', 'hard_or_soft', 'soft_default'):
                if field == 'area_sqm' or field.startswith('mrt_'):
                    reason = 'range_or_route_check_required'
                elif 'conflicting_' + field in flags or (field == 'gender_restriction' and 'gender_label_text_conflict' in flags):
                    reason = 'conflicting_evidence'
                else:
                    state = 'pass' if value == wanted else 'fail'
                    reason = 'explicit_value_comparison'
        checks.append(dict(requirement=key, value=value, requested=wanted, result=state, reason=reason))
    status = 'fail' if any(c['result'] == 'fail' for c in checks) else 'needs_confirmation' if any(c['result'] == 'needs_confirmation' for c in checks) else 'pass'
    return {'status': status, 'checks': checks}


def sheet(wb, name, records, fields=None):
    ws = wb.create_sheet(name)
    fields = fields or list(records[0])
    ws.append(fields)
    for record in records:
        ws.append([json.dumps(record.get(f), ensure_ascii=False) if isinstance(record.get(f), (list, dict)) else record.get(f) for f in fields])
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = 's'
            if fields[cell.column - 1] == 'listing_id':
                cell.number_format = '@'
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='234E70')
        ws.column_dimensions[cell.column_letter].width = 24
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = ws.dimensions
    return ws


def main():
    input_path = ROOT / 'output-600/listings_600.json'
    before = hashlib.sha256(input_path.read_bytes()).hexdigest()
    raw = json.loads(input_path.read_text(encoding='utf-8'))
    cards = rendered_cards({str(r['listing_id']) for r in raw})
    rows, provenance, duplicate_records = [], [], []
    seen = {}
    for source in raw:
        key = (str(source['listing_id']), canonical_url(source['url']))
        if key in seen and source == seen[key]:
            duplicate_records.append({'listing_id': key[0], 'reason': 'exact_duplicate'})
            continue
        seen[key] = source
        row, prov = clean_record(source, cards[str(source['listing_id'])])
        rows.append(row)
        provenance.append(prov)
    for key in ('listing_id', 'listing_url', 'address'):
        groups = defaultdict(list)
        for r in rows:
            if not missing(r[key]):
                groups[text(r[key]).casefold()].append(r)
        for members in groups.values():
            if len(members) > 1:
                flag = 'same_address_different_ids' if key == 'address' else 'duplicate_' + key
                for r in members:
                    r['review_flags'] = sorted(set(r['review_flags'] + [flag]))
    # Keep the full field contract in the dictionary, but omit zero-evidence columns
    # from the working view. Missing keys must still be treated as unknown by clients.
    deferred = [f for f in META if f != 'review_flags' and all(missing(r[f]) for r in rows)]
    fields = [f for f in META if f not in deferred]
    exported = [{f: r[f] for f in fields} for r in rows]
    dictionary = []
    for f, meta in META.items():
        valid = sum(not missing(r[f]) and r[f] != [] for r in rows)
        dictionary.append(dict(meta, valid_count=valid, unknown_or_blank_count=len(rows)-valid,
                               coverage=round(valid/len(rows), 4),
                               output_status='deferred_no_evidence' if f in deferred else 'retained',
                               missing_policy='empty_list' if f == 'review_flags' else 'null' if f in ('rent_monthly_sgd', 'bedrooms', 'bathrooms', 'area_sqm', 'mrt_distance_m', 'mrt_minutes_reported', 'available_from', 'listed_date', 'collected_at', 'tenant_count_reported', 'minimum_lease_months') else 'unknown'))
    missing_audit = []
    for f in fields:
        null_count = sum(r[f] is None for r in rows)
        unknown_count = sum(r[f] == UNKNOWN for r in rows)
        empty_list_count = sum(r[f] == [] for r in rows)
        missing_audit.append(dict(field=f, row_count=len(rows), null_count=null_count,
                                  unknown_count=unknown_count, empty_list_count=empty_list_count,
                                  populated_count=len(rows)-null_count-unknown_count-empty_list_count,
                                  excel_display='blank_means_missing_number_or_date' if null_count else
                                  'empty_list_means_no_flags' if f == 'review_flags' else 'unknown_means_unspecified'))
    flag_counts = Counter(f for r in rows for f in r['review_flags'])
    summary = dict(input_rows=len(raw), output_rows=len(rows), exported_fields=len(fields),
                   exact_duplicates_removed=len(duplicate_records), deferred_fields=deferred,
                   input_sha256=before, district_counts=dict(Counter(r['district_code'] for r in rows)),
                   rental_scope_counts=dict(Counter(r['rental_scope'] for r in rows)),
                   restriction_coverage={f: sum(not missing(r[f]) for r in rows) for f in RESTRICTIONS},
                   flags=dict(flag_counts), flagged_rows=sum(bool(r['review_flags']) for r in rows),
                   transformations=dict(Counter(t for p in provenance for t in p['transformations'])))
    OUT.mkdir(exist_ok=True)
    for name, obj in [('agent_listings.json', exported), ('field_dictionary.json', dictionary),
                      ('provenance.json', provenance), ('cleaning_report.json', summary),
                      ('missing_value_audit.json', missing_audit), ('removed_duplicates.json', duplicate_records)]:
        (OUT / name).write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding='utf-8')
    with (OUT / 'agent_listings.csv').open('w', encoding='utf-8-sig', newline='') as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows({k: json.dumps(v, ensure_ascii=False) if isinstance(v, list) else v for k, v in r.items()} for r in exported)
    wb = Workbook()
    wb.remove(wb.active)
    sheet(wb, 'Agent_Listings', exported)
    sheet(wb, 'Field_Dictionary', dictionary)
    sheet(wb, 'Missing_Value_Audit', missing_audit)
    sheet(wb, 'Evidence', provenance)
    sheet(wb, 'Review_Issues', [dict(listing_id=r['listing_id'], review_flags=r['review_flags']) for r in rows if r['review_flags']], ['listing_id', 'review_flags'])
    sheet(wb, 'Cleaning_Summary', [{'metric': k, 'value': v} for k, v in summary.items()])
    wb.save(OUT / 'nus_rental_agent_cleaned_600_v2.xlsx')
    write_guide(summary, dictionary, raw[0].keys(), fields)
    assert hashlib.sha256(input_path.read_bytes()).hexdigest() == before
    check = load_workbook(OUT / 'nus_rental_agent_cleaned_600_v2.xlsx', read_only=True)
    assert check['Agent_Listings'].max_row == len(rows) + 1
    assert all(isinstance(row[0].value, str) for row in check['Agent_Listings'].iter_rows(min_row=2))
    check.close()
    print(json.dumps(summary, indent=2))


def write_guide(summary, dictionary, old_fields, fields):
    renamed = {'url': 'listing_url', 'price_sgd': 'rent_monthly_sgd', 'price_terms': 'rent_price_type',
               'posted_at_text': 'listed_date', 'floor_area_sqm': 'area_sqm', 'floor_area_sqft': 'area_sqm',
               'land_area_sqft': 'area_sqm (fallback only; raw retained)', 'area_sqft': 'area_sqm (fallback only)',
               'availability_text': 'available_from + availability_status',
               'male_preference_in_card_description': 'gender_preference (explicit headline evidence)'}
    archived = {'source_url', 'sample_cohort', 'card_text', 'thumbnail_url', 'image_urls',
                'field_evidence', 'tenant_label_text', 'tenant_label_source_files'}
    actions = []
    for old in old_fields:
        action = 'retained_and_normalized' if old in fields else 'merged_or_renamed' if old in renamed else 'moved_to_evidence' if old in archived else 'excluded_from_agent_view_raw_preserved'
        actions.append({'original_field': old, 'action': action, 'destination': renamed.get(old, old if old in fields else 'provenance.json' if old in archived else 'output-600/listings_600.json')})
    (OUT / 'field_actions.json').write_text(json.dumps(actions, indent=2), encoding='utf-8')
    dictionary_wb = Workbook()
    dictionary_wb.remove(dictionary_wb.active)
    sheet(dictionary_wb, 'Field_Dictionary', dictionary)
    sheet(dictionary_wb, 'Original_Field_Actions', actions)
    dictionary_wb.save(OUT / 'agent_field_dictionary_bilingual_v2.xlsx')
    lines = [
        '# NUS Rental Agent：第一版清洗说明', '',
        f"输入 600 条，输出 {summary['output_rows']} 条；主表 {summary['exported_fields']} 个英文字段。原始 500 条和补充 100 条均保留，优先七区仍为 450 条，其余区域 150 条。",
        '这是保存网页的历史快照，不是实时可租库存，也不是全市场随机样本。', '',
        '## 从哪里开始', '',
        '- 小组评审：打开 `nus_rental_agent_cleaned_600_v2.xlsx`，先看 `Agent_Listings`，有问题再看 `Review_Issues` 和 `Evidence`。',
        '- 字段解释：打开 `agent_field_dictionary_bilingual_v2.xlsx`，含英文名、中文含义、有效数量、缺失数、使用角色、旧字段处理去向。',
        '- Agent 接入：使用 `agent_listings.json`；CSV 供表格交换，Excel 中 ID 已存为文本。',
        '- 原始文件保留在 `output-600`，此次不覆盖；运行 `python prepare_agent_data.py` 可重建本目录输出。', '',
        '## Hard constraints 与 Soft preferences', '',
        '默认硬条件：`rent_monthly_sgd`（结合 `rent_price_type`）、`rental_scope`、`bedrooms`。预算仅比较广告月租；没有自动分摊租金，也没有把未给出的水电、押金、中介费算成 0。',
        '默认软偏好：`property_type`、`room_type`、`area_sqm`（面积范围必须可比）、`furnishing`、`mrt_distance_m`、`mrt_minutes_reported`、`district_code`、`landlord_lives_in`。',
        '条件角色由用户决定：例如“必须能做饭”使 `cooking_policy` 成为硬条件；“最好能做饭”才是软偏好。“只能 D05”“必须独立卫浴”“必须允许宠物”同理。',
        '`gender_restriction` 是明确排他限制；`gender_preference` 是偏好，不能作为性别排除规则。平台 `tenant_gender_label` 单独保留，只能据此提示核实，不代表已有室友性别。不要推断用户性别。',
        '`tenant_count_reported` 的适用范围未确认，不能直接当作已入住人数或可接受人数；`tenant_count_scope` 暂为 unknown。',
        'NUS 通勤、坐标、家具细项、完整账单费用等没有可靠数据，目前不能给出已满足判断；MRT 分钟数不能冒充到 NUS 的通勤分钟数。车站未写 U/C 也不保证已运营。', '',
        '## Unknown 的执行规则', '',
        '- 数值和精确日期缺失：JSON null / Excel 空白。分类缺失：小写 `unknown`。不填 0、不填均值，不让 LLM 从常识补齐。',
        '- Excel 数值单元格保持空白，避免把数字列变成文字列；界面可以显示“未注明”，但不要把这几个字写回数值数据。`Missing_Value_Audit` 分开统计 null、unknown 和空列表。',
        '- 全部为 null/unknown 的字段移出主表，字段字典标为 deferred_no_evidence。`available_from`、`tenant_count_scope` 当前都没有有效值；`listed_date` 当前 600 条有效，继续保留。',
        '- 硬条件使用 pass / fail / needs_confirmation 三态。未知不能通过；另列待确认候选，不自动放宽条件。',
        '- 软偏好缺失不等于差：建议已知匹配度与证据覆盖率分开呈现，不能只凭少量已知字段给出高置信推荐。',
        '- 主表未输出的字段也必须按 unknown/null 读取；“没有这一列”绝不等于“没有限制”。',
        '- `prepare_agent_data.assess_constraints` 提供基础三态接口，支持预算、最少卧室数及明确分类条件；通勤、日期范围、入住人数等复杂条件返回待确认，未实现的比较不能默默通过。', '',
        '## 实际处理', '',
        '1. 按房源 ID 和规范化详情 URL 检查重复；当前 600 个 ID 和 URL 唯一，未删除房源。同地址不同编号只标记，不假设同一套房。',
        '2. 月租和面积转数值；平方米 = sqft × 0.09290304，优先已报告 floor 面积。土地面积保留在原数据，不与楼面面积相加。租房间时展示面积未必是房间面积，因此范围保留 unknown。',
        '3. 房型和物业类型统一英文枚举。Studio 只有在渲染后的卧室标签明确写出时才记 0 个独立卧室；不能因此断言整租。',
        '4. 非正数卫生间数置空，保留原值和修改记录；不认为房源没有卫生间。租金和面积非正值置空并标记。',
        '5. 发布日期统一 YYYY-MM-DD，采集时间统一 ISO。可入住时间没有年份则精确日期留空，并标记 date_missing_year；不根据今天猜年份。Ready to Move 仅是快照时的声明。',
        '6. 限制提取使用保存网页对应房源的广告 headline，精确规则抽取并附原文证据；不使用整页其它房源文本、不使用 LLM 猜测。规则保守，未匹配不表示全文肯定不存在该信息。',
        '7. 整租只接受明确 whole/entire unit 等文字；有房间标签或明确 common/master/shared room 才归房间。卧室数不会自动转换为整租判断。',
        '8. 未带价格修饰词的广告报价记 asking，仅表示广告报价；不承诺可成交。Starting From/Guide Price 需要确认实际价格。',
        '9. 低于 S$300 或高于 S$30,000 月租、低于 80 或高于 10,000 sqft 面积暂列复核；这些是项目审阅阈值，不是市场真伪判定。高价大房子保留，不自动截尾或删除。',
        '10. 比较广告单价与月租/面积（楼面与土地面积均无法解释且偏差超过 15% 才标记；PPG 有些单价基于土地面积）、原平方米与换算值（超过 2% 且大于 1 平方米）、人数标题和标签、房型范围矛盾。来源矛盾保留证据。',
        '11. 重复性元数据、经纪人资料、项目编号、产权年限、调试字段不进 Agent 决策表；图片和广告原文放 Evidence，仍可用于展示和核实。旧 review_flags 保留在 Evidence，新表重算，已修复问题不会当作仍未修复重复展示。', '',
        '## 统计', '',
        '```json', json.dumps(summary, ensure_ascii=False, indent=2), '```', '',
        '## 主表全部字段', '', '| Field | 中文含义 | 有效数量 | 默认用途 |', '|---|---|---:|---|',
    ]
    lines += [f"| {d['field']} | {d['meaning_zh']} | {d['valid_count']} | {d['decision_role']} |" for d in dictionary if d['output_status'] == 'retained']
    lines += ['', '## 暂不展示但保留在字段契约中的字段', '',
              '以下字段未提取到可靠值，主表先不展示。后续用户问到时必须返回 unknown，而不是 no：', '',
              ', '.join(summary['deferred_fields']), '',
              '另需后续补采：NUS 目的地及真实通勤路线/时间、邮编及坐标、押金、中介费、总月度费用、租期和目前可租状态。未给出这些数据前不能据此做确定性推荐。']
    (OUT / 'CLEANING_REPORT_ZH.md').write_text('\n'.join(lines), encoding='utf-8')


if __name__ == '__main__':
    main()
